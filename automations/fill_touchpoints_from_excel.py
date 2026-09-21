"""Fill missing touchpoint fields on USEUContact rows from an Excel file.

Usage (on server, from the automations/ directory):
    /var/www/Magnum\ Opus\ Consultants/Automation-Platform/venv/bin/python3 fill_touchpoints_from_excel.py
    # or with an explicit path:
    python3 fill_touchpoints_from_excel.py --excel "Book1 4.xlsx"
    # dry-run is the default. To actually write:
    python3 fill_touchpoints_from_excel.py --apply

Rules:
    - Match an Excel row to a DB row only when org_name, contact_name, phone, and
      email all match exactly (case-insensitive, whitespace-trimmed).
    - For each matched DB row, only fill a field if the DB field is currently empty
      AND the Excel has a value for it. NEVER overwrites. NEVER deletes.
    - If an Excel row matches multiple DB rows, fill each matching row (same rule
      per row: only blanks get filled).
    - Excel rows with no DB match are reported but not created.
"""
import argparse
import os
import sys
from collections import Counter

import django
import openpyxl

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'automations.settings')
django.setup()

from dashboard.models import USEUContact  # noqa: E402


EXCEL_TO_DB_FIELDS = [
    ('Touchpoint 1', 'touchpoint_1'),
    ('TP1 Sent On', 'tp1_sent_on'),
    ('Touchpoint 2', 'touchpoint_2'),
    ('Last Touch Date', 'last_touch'),
    ('Journey Status', 'journey_status'),
    ('TP1ProcessingId', 'tp1_processing_id'),
]

MATCH_FIELDS = [
    ('Org. Name', 'org_name'),
    ('Contact Name', 'contact_name'),
    ('Phone', 'phone'),
    ('Email', 'email'),
]


def norm(v):
    if v is None:
        return ''
    return str(v).strip().lower()


def read_excel(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    rows = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        it = ws.iter_rows(values_only=True)
        try:
            header = next(it)
        except StopIteration:
            continue
        header = [str(h).strip() if h is not None else '' for h in header]
        idx = {h: i for i, h in enumerate(header)}
        required = [c for c, _ in MATCH_FIELDS] + [c for c, _ in EXCEL_TO_DB_FIELDS]
        missing = [c for c in required if c not in idx]
        if missing:
            print(f'[skip sheet "{sheet_name}"] missing columns: {missing}')
            continue
        for row in it:
            if row is None:
                continue
            rec = {}
            for col in required:
                i = idx[col]
                rec[col] = row[i] if i < len(row) else None
            rows.append(rec)
    return rows


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--excel', default='Book1 4.xlsx', help='Path to the Excel file')
    parser.add_argument('--apply', action='store_true', help='Actually write changes (default is dry-run)')
    parser.add_argument('--verbose', action='store_true', help='Print each row change')
    args = parser.parse_args()

    if not os.path.exists(args.excel):
        print(f'ERROR: Excel file not found: {args.excel}')
        sys.exit(1)

    print(f'Reading {args.excel} ...')
    excel_rows = read_excel(args.excel)
    print(f'Excel rows: {len(excel_rows)}')

    # Index DB by (org, contact, phone, email) normalized
    print('Indexing DB (useu_contacts) ...')
    db_index = {}
    total_db = 0
    for c in USEUContact.objects.all().iterator():
        total_db += 1
        key = (norm(c.org_name), norm(c.contact_name), norm(c.phone), norm(c.email))
        db_index.setdefault(key, []).append(c)
    print(f'DB rows: {total_db}')

    stats = Counter()
    field_fills = Counter()
    unmatched_excel = []
    to_save = {}  # id -> (contact, [field_names])

    for rec in excel_rows:
        key = tuple(norm(rec[col]) for col, _ in MATCH_FIELDS)
        matches = db_index.get(key, [])
        if not matches:
            stats['excel_no_match'] += 1
            unmatched_excel.append(key)
            continue
        stats['excel_matched'] += 1
        stats['db_rows_matched'] += len(matches)

        for contact in matches:
            fields_to_fill = []
            for excel_col, db_field in EXCEL_TO_DB_FIELDS:
                excel_val = rec.get(excel_col)
                if excel_val is None or str(excel_val).strip() == '':
                    continue
                current = getattr(contact, db_field, '') or ''
                if str(current).strip() != '':
                    continue  # never overwrite
                fields_to_fill.append((db_field, str(excel_val).strip()))
                field_fills[db_field] += 1
            if fields_to_fill:
                stats['db_rows_would_change'] += 1
                to_save[contact.id] = (contact, fields_to_fill)
                if args.verbose:
                    print(f'  [{contact.id}] {contact.org_name} | {contact.email} -> {[f for f, _ in fields_to_fill]}')

    print('\n===== DRY-RUN REPORT =====' if not args.apply else '\n===== APPLY REPORT =====')
    print(f'Excel rows total:               {len(excel_rows)}')
    print(f'Excel rows matched to DB:       {stats["excel_matched"]}')
    print(f'Excel rows with no DB match:    {stats["excel_no_match"]}')
    print(f'DB rows matched (total):        {stats["db_rows_matched"]}')
    print(f'DB rows that would change:      {stats["db_rows_would_change"]}')
    print('Field-by-field fills:')
    for f, n in field_fills.most_common():
        print(f'  {f}: {n}')

    if unmatched_excel:
        print(f'\nFirst 10 unmatched Excel rows (org, contact, phone, email):')
        for u in unmatched_excel[:10]:
            print(f'  {u}')

    if not args.apply:
        print('\nThis was a DRY-RUN. No changes written. Re-run with --apply to write.')
        return

    print('\nApplying changes ...')
    written = 0
    for _id, (contact, fields) in to_save.items():
        for db_field, val in fields:
            setattr(contact, db_field, val)
        contact.save(update_fields=[f for f, _ in fields])
        written += 1
    print(f'Done. Wrote {written} rows.')


if __name__ == '__main__':
    main()
