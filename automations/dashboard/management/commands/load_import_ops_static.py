"""One-off loader for the two static Import Operations reference tables.

Loads:
  - Headers.xlsx       → import_operations_headers       (1 row of label text)
  - New Source.xlsx    → import_operations_new_source    (historical shipments)

Both use the same 121-column schema as import_ops. These are STATIC — run once.

Usage:
    py manage.py load_import_ops_static \\
        --headers /path/to/Headers.xlsx \\
        --new-source /path/to/'New Source.xlsx'
"""
import io

import openpyxl
from django.core.management.base import BaseCommand
from django.db import connection, transaction
from psycopg2.extras import execute_values

from dashboard.scheduler import IMPORT_OPS_COLUMNS, _coerce_io_value


def _build_create_ddl(table):
    cols_ddl = ',\n    '.join(f'{n} {t}' for n, t in IMPORT_OPS_COLUMNS)
    return f'''
    DROP TABLE IF EXISTS {table};
    CREATE TABLE {table} (
        id BIGSERIAL PRIMARY KEY,
        {cols_ddl}
    );
    '''


def _parse_shipment_profile(path, force_text=False):
    """Read a Shipment Profile sheet where row 1 is the header row, data starts row 2.
    If force_text, every value is stored as text regardless of the column type
    (used for Headers.xlsx where the 'data' is the header labels themselves)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb['Shipment Profile']
    rows_out = []
    n_cols = len(IMPORT_OPS_COLUMNS)
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        vals = list(row[:n_cols])
        if len(vals) < n_cols:
            vals += [None] * (n_cols - len(vals))
        if not any(v not in (None, '') for v in vals):
            continue
        if force_text:
            coerced = tuple(str(v).strip() if v is not None else None for v in vals)
        else:
            coerced = tuple(_coerce_io_value(t, v) for (_, t), v in zip(IMPORT_OPS_COLUMNS, vals))
        rows_out.append(coerced)
    wb.close()
    return rows_out


class Command(BaseCommand):
    help = 'Load Headers.xlsx + New Source.xlsx into static Import Operations tables.'

    def add_arguments(self, parser):
        parser.add_argument('--headers', required=True,
                            help='Path to Headers.xlsx')
        parser.add_argument('--new-source', required=True,
                            help='Path to New Source.xlsx')

    def handle(self, *args, **opts):
        col_names = [n for n, _ in IMPORT_OPS_COLUMNS]
        insert_cols = ','.join(col_names)

        # ── Headers ─────────────────────────────────────────────
        self.stdout.write(f'Loading Headers from {opts["headers"]}')
        hdr_rows = _parse_shipment_profile(opts['headers'], force_text=True)
        if not hdr_rows:
            # File has only the header row (r1) and no data rows — load row 1
            # verbatim as the single data row.
            wb = openpyxl.load_workbook(opts['headers'], read_only=True, data_only=True)
            r1 = next(wb['Shipment Profile'].iter_rows(min_row=1, max_row=1, values_only=True))
            vals = list(r1[:len(IMPORT_OPS_COLUMNS)])
            if len(vals) < len(IMPORT_OPS_COLUMNS):
                vals += [None] * (len(IMPORT_OPS_COLUMNS) - len(vals))
            hdr_rows = [tuple(str(v).strip() if v is not None else None for v in vals)]
            wb.close()

        with transaction.atomic():
            with connection.cursor() as cur:
                # Headers table — all text, since it holds label strings
                cur.execute('DROP TABLE IF EXISTS import_operations_headers')
                cols_text_ddl = ',\n    '.join(f'{n} TEXT' for n in col_names)
                cur.execute(f'CREATE TABLE import_operations_headers (id BIGSERIAL PRIMARY KEY, {cols_text_ddl})')
                execute_values(cur,
                    f'INSERT INTO import_operations_headers ({insert_cols}) VALUES %s',
                    hdr_rows)
        self.stdout.write(self.style.SUCCESS(
            f'  → import_operations_headers loaded ({len(hdr_rows)} row, {len(col_names)} cols)'))

        # ── New Source ──────────────────────────────────────────
        self.stdout.write(f'\nLoading New Source from {opts["new_source"]}')
        src_rows = _parse_shipment_profile(opts['new_source'])
        with transaction.atomic():
            with connection.cursor() as cur:
                cur.execute(_build_create_ddl('import_operations_new_source'))
                cur.execute('CREATE INDEX import_operations_new_source_shipment_id_idx '
                            'ON import_operations_new_source (shipment_id)')
                execute_values(cur,
                    f'INSERT INTO import_operations_new_source ({insert_cols}) VALUES %s',
                    src_rows)
        self.stdout.write(self.style.SUCCESS(
            f'  → import_operations_new_source loaded ({len(src_rows):,} rows, {len(col_names)} cols)'))

        # ── Summary ──────────────────────────────────────────────
        with connection.cursor() as cur:
            cur.execute('SELECT COUNT(*) FROM import_operations_headers')
            h_count = cur.fetchone()[0]
            cur.execute('SELECT COUNT(*) FROM import_operations_new_source')
            n_count = cur.fetchone()[0]
        self.stdout.write(f'\nDone. headers={h_count}  new_source={n_count:,}')
