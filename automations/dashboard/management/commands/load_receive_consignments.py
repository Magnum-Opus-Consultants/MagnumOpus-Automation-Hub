"""Load Receive Consignments exports into the database.

    python manage.py load_receive_consignments --folder "<path to exports>"

The Bruce report used to read the .xlsx files straight from a folder and work
everything out in Power Query. This does that work once, on the way in, so the
report becomes a plain read of one table and every other consumer - a query, a
page on the platform - sees the same numbers.

Re-running is safe. A file already loaded, unchanged, is skipped; pass --force
to load it again. After every run the latest-snapshot flags are recomputed
across the whole table, because "latest" is a property of the set, not of the
file that happened to arrive last.
"""
import datetime as dt
import os
import re

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone

from dashboard.models import ReceiveConsignment, ReceiveConsignmentImport

# CargoWise names its exports Export_Receiveconsignments - 2026-08-17T201847.010.xlsx
_TS_IN_NAME = re.compile(r'(\d{4})-(\d{2})-(\d{2})T(\d{2})(\d{2})(\d{2})')

# Sheet columns, as exported, mapped to model fields.
COLUMNS = {
    'Next open service': 'next_open_service',
    'Receive consignment ID': 'receive_consignment_id',
    'RCN reference': 'rcn_reference',
    'Consignor': 'consignor',
    'Consignee': 'consignee',
    'Number of packages': 'number_of_packages',
    'BKD': 'bkd',
    'ARV': 'arv',
    'CTT': 'ctt',
    'PIC': 'pic',
    'PUT': 'put',
    'In warehouse': 'in_warehouse',
    'Warehouse': 'warehouse',
    'Completion date': 'completion_date',
    'Next discharge port': 'next_discharge_port',
    'Service level': 'service_level',
    'Overs': 'overs',
    'Booking party': 'booking_party',
}
_INTS = {'number_of_packages', 'bkd', 'arv', 'ctt', 'pic', 'put',
         'in_warehouse', 'overs'}
_TEXT_MAX = {'receive_consignment_id': 60, 'rcn_reference': 60,
             'next_open_service': 120, 'warehouse': 120,
             'service_level': 120, 'next_discharge_port': 120}


def party_name(party):
    """Trading name from a CargoWise address string.

    They are written "NAME, STREET, CITY STATE ZIP, COUNTRY", so the last three
    comma-separated segments are the address and anything before them is the
    name. Short strings fall back to the first segment. Same rule the report's
    own fnPartyName used, kept so figures do not move in the changeover.
    """
    if not party or not party.strip():
        return ''
    parts = [p.strip() for p in party.split(',')]
    name = ', '.join(parts[:-3]) if len(parts) >= 4 else parts[0]
    return (name or parts[0]).strip()


def party_country(party):
    """Country - the final segment - from a CargoWise address string."""
    if not party or not party.strip():
        return ''
    parts = [p.strip() for p in party.split(',')]
    return parts[-1] if len(parts) >= 2 and parts[-1] else ''


def split_code(value):
    """Split a CargoWise "CODE - Description" into its two halves.

    Only the first dash separates them - warehouse names contain dashes of
    their own - and a value with no dash is treated as all name, since that is
    what a reader would call it.
    """
    if not value or not value.strip():
        return '', ''
    parts = value.split(' - ', 1)
    if len(parts) == 2:
        return parts[0].strip(), parts[1].strip()
    return '', value.strip()


def age_bucket(days):
    """The report's buckets, with their sort order alongside.

    Returned together because a bucket and its position are one decision; split
    across two functions they drift apart the first time a band is edited.
    """
    if days is None:
        return 'Unknown', 99
    if days <= 7:
        return '0-7 days', 1
    if days <= 14:
        return '8-14 days', 2
    if days <= 30:
        return '15-30 days', 3
    if days <= 60:
        return '31-60 days', 4
    return '60+ days', 5


def timestamp_from_name(filename):
    """The export's own moment, read out of its filename."""
    m = _TS_IN_NAME.search(filename)
    if not m:
        return None
    y, mo, d, h, mi, s = (int(x) for x in m.groups())
    try:
        naive = dt.datetime(y, mo, d, h, mi, s)
    except ValueError:
        return None
    return timezone.make_aware(naive, timezone.get_default_timezone())


def _as_int(v):
    if v is None or v == '':
        return 0
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return 0


def _as_dt(v):
    if isinstance(v, dt.datetime):
        return timezone.make_aware(v, timezone.get_default_timezone()) \
            if timezone.is_naive(v) else v
    return None


class Command(BaseCommand):
    help = 'Load Receive Consignments .xlsx exports into the database.'

    def add_arguments(self, parser):
        parser.add_argument('--folder', required=True,
                            help='Folder holding Export_Receiveconsignments*.xlsx')
        parser.add_argument('--force', action='store_true',
                            help='Reload files that have already been loaded.')

    def handle(self, *args, **opts):
        try:
            import openpyxl
        except ImportError:
            raise CommandError('openpyxl is required: pip install openpyxl')

        folder = opts['folder']
        if not os.path.isdir(folder):
            raise CommandError(f'Not a folder: {folder}')

        files = sorted(f for f in os.listdir(folder)
                       if f.lower().startswith('export_receiveconsignments')
                       and f.lower().endswith(('.xlsx', '.xlsm'))
                       and not f.startswith('~$'))
        if not files:
            raise CommandError(f'No Export_Receiveconsignments*.xlsx in {folder}')

        loaded = skipped = 0
        for name in files:
            path = os.path.join(folder, name)
            size = os.path.getsize(path)
            modified = timezone.make_aware(
                dt.datetime.fromtimestamp(os.path.getmtime(path)),
                timezone.get_default_timezone())

            existing = ReceiveConsignmentImport.objects.filter(
                filename=name, file_size=size).first()
            if existing and not opts['force']:
                self.stdout.write(f'  skip   {name} (already loaded)')
                skipped += 1
                continue
            if existing:
                existing.delete()   # cascades its rows; --force means replace

            n = self._load_file(openpyxl, path, name, size, modified)
            self.stdout.write(self.style.SUCCESS(f'  loaded {name} ({n} rows)'))
            loaded += 1

        self._mark_latest()
        total = ReceiveConsignment.objects.count()
        unbooked = ReceiveConsignment.objects.filter(
            is_unbooked=True, is_latest_snapshot=True, in_warehouse__gt=0).count()
        self.stdout.write('')
        self.stdout.write(self.style.SUCCESS(
            f'{loaded} file(s) loaded, {skipped} skipped. '
            f'{total} rows in table; {unbooked} unbooked awaiting action.'))

    @transaction.atomic
    def _load_file(self, openpyxl, path, name, size, modified):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else '' for h in next(rows)]

        exported = timestamp_from_name(name)
        imp = ReceiveConsignmentImport.objects.create(
            filename=name, file_size=size, file_modified=modified,
            export_timestamp=exported)

        today = timezone.localdate()
        batch = []
        for raw in rows:
            if raw is None or all(v is None or v == '' for v in raw):
                continue
            rec = dict(zip(header, raw))
            vals = {}
            for col, field in COLUMNS.items():
                v = rec.get(col)
                if field in _INTS:
                    vals[field] = _as_int(v)
                elif field == 'completion_date':
                    vals[field] = _as_dt(v)
                else:
                    s = '' if v is None else str(v).strip()
                    cap = _TEXT_MAX.get(field)
                    vals[field] = s[:cap] if cap else s

            if not vals.get('receive_consignment_id'):
                continue

            completion = vals.get('completion_date')
            days = (today - timezone.localtime(completion).date()).days \
                if completion else None
            bucket, order = age_bucket(days)
            booking = vals.get('booking_party', '')

            consignee = vals.get('consignee', '')
            wh_code, wh_name = split_code(vals.get('warehouse', ''))
            sl_code, sl_name = split_code(vals.get('service_level', ''))

            batch.append(ReceiveConsignment(
                source=imp,
                consignor_name=party_name(vals.get('consignor', ''))[:200],
                consignor_country=party_country(vals.get('consignor', ''))[:80],
                consignee_name=party_name(consignee)[:200],
                consignee_country=party_country(consignee)[:80],
                booking_party_name=party_name(booking)[:200],
                booking_party_country=party_country(booking)[:80],
                warehouse_code=wh_code[:40], warehouse_name=wh_name[:120],
                service_level_code=sl_code[:40], service_level_name=sl_name[:120],
                # The grid's own test: the booking party literally reads UNBOOKED.
                is_unbooked='UNBOOKED' in booking.upper(),
                has_overs=vals.get('overs', 0) > 0,
                has_consignee=bool(consignee),
                has_discharge_port=bool(vals.get('next_discharge_port')),
                has_next_open_service=bool(vals.get('next_open_service')),
                age_days=days, age_bucket=bucket, age_bucket_order=order,
                completion_date_only=(timezone.localtime(completion).date()
                                      if completion else None),
                source_file=name, export_timestamp=exported,
                file_modified=modified,
                row_key=f'{name}|{vals["receive_consignment_id"]}'[:480],
                **vals))

        ReceiveConsignment.objects.bulk_create(batch, batch_size=500)
        imp.row_count = len(batch)
        imp.save(update_fields=['row_count'])
        wb.close()
        return len(batch)

    def _mark_latest(self):
        """Flag only the newest appearance of each consignment.

        Exports overlap, so the same consignment arrives in several files. Left
        unflagged every package count is multiplied by however many files
        mention it. Recomputed over the whole table, not just what was loaded,
        because an older file loaded late still must not win.
        """
        newest = {}
        for pk, rcid, ts in ReceiveConsignment.objects.values_list(
                'id', 'receive_consignment_id', 'export_timestamp').iterator():
            key = rcid
            cur = newest.get(key)
            # No timestamp sorts oldest, so a named export always beats one
            # whose filename did not carry a moment.
            rank = (ts or dt.datetime.min.replace(tzinfo=dt.timezone.utc), pk)
            if cur is None or rank > cur[1]:
                newest[key] = (pk, rank)

        keep = {pk for pk, _ in newest.values()}
        ReceiveConsignment.objects.exclude(id__in=keep).update(is_latest_snapshot=False)
        ReceiveConsignment.objects.filter(id__in=keep).update(is_latest_snapshot=True)
