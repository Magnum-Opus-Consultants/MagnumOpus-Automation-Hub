import os, django, base64, io, re, time, requests
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'automations.settings')
django.setup()
from decimal import Decimal
from collections import defaultdict
import openpyxl, xlrd
from django.db import connection
from psycopg2.extras import execute_values
from dashboard.scheduler import EMAIL_MAILBOX
from dashboard.views import _get_graph_token


class Tok:
    def __init__(self): self.t = _get_graph_token(); self.ts = time.time()
    def get(self):
        if time.time() - self.ts > 2400:
            self.t = _get_graph_token(); self.ts = time.time()
        return self.t


T = Tok()


def req(url, timeout=60, retries=4):
    for i in range(retries):
        try:
            r = requests.get(url, headers={'Authorization': f'Bearer {T.get()}'}, timeout=timeout)
            if r.status_code == 401:
                T.t = _get_graph_token(); T.ts = time.time(); continue
            if r.status_code == 429:
                time.sleep(int(r.headers.get('Retry-After', '10'))); continue
            r.raise_for_status(); return r
        except (requests.ConnectionError, requests.Timeout):
            if i == retries - 1: raise
            time.sleep(5 * (i + 1))


def scan_all():
    url = (f'https://graph.microsoft.com/v1.0/users/{EMAIL_MAILBOX}/messages'
           f'?$top=100&$orderby=receivedDateTime desc'
           f'&$select=id,subject,receivedDateTime,from,hasAttachments')
    out = []
    while url:
        r = req(url, 30); j = r.json()
        out.extend(j.get('value', []))
        url = j.get('@odata.nextLink')
    return out


def get_xlsx(msg_id):
    ar = req(f'https://graph.microsoft.com/v1.0/users/{EMAIL_MAILBOX}/messages/{msg_id}/attachments', 60)
    x = next((a for a in ar.json().get('value', [])
              if a.get('name', '').lower().endswith(('.xlsx', '.xls'))), None)
    return (x['name'], base64.b64decode(x['contentBytes'])) if x else (None, None)


print('Scanning inbox...')
all_msgs = scan_all()
print(f'{len(all_msgs)} total')

# ========== IMPORT OPERATIONS ==========
io_matches = [m for m in all_msgs
              if 'import operational report moc' in m.get('subject', '').lower()
              and m.get('hasAttachments')]
io_matches.sort(key=lambda m: m['receivedDateTime'], reverse=True)
print(f'IMPORT OPS: {len(io_matches)} emails; latest={io_matches[0]["receivedDateTime"][:10] if io_matches else "-"}')
if io_matches:
    fname, fb = get_xlsx(io_matches[0]['id'])
    wb = openpyxl.load_workbook(io.BytesIO(fb), read_only=True, data_only=True)
    ws = wb['Shipment Profile']
    rows = []
    for row in ws.iter_rows(min_row=15, values_only=False):
        cells = [c.value for c in row]
        vals = cells[2:25] if len(cells) > 24 else cells[2:] + [None] * (23 - max(0, len(cells) - 2))
        if not vals[0]: continue
        rows.append(tuple(str(v).strip() if v is not None else '' for v in vals))
    wb.close()
    with connection.cursor() as cur:
        cur.execute('DELETE FROM import_ops')
        execute_values(cur,
            'INSERT INTO import_ops (shipment_id,shipment_direction,report_date,trans,customs_info,mode,origin,origin_country,destination,destination_country,consignor_code,consignor_name,consignee_code,consignee_name,house_ref,incoterm,additional_terms,ppd_ccx,goods_description,origin_etd,destination_eta,weight,weight_unit) VALUES %s',
            rows)
    print(f'  -> inserted {len(rows)} import_ops rows')

# ========== WIP & ACCRUAL ==========
wip_matches = [m for m in all_msgs
               if 'wip rev and accrued costs' in m.get('subject', '').lower()
               and m.get('hasAttachments')]
wip_matches.sort(key=lambda m: m['receivedDateTime'], reverse=True)
print(f'WIP: {len(wip_matches)} emails; latest={wip_matches[0]["receivedDateTime"][:10] if wip_matches else "-"}')
if wip_matches:
    fname, fb = get_xlsx(wip_matches[0]['id'])
    wb = openpyxl.load_workbook(io.BytesIO(fb), read_only=True, data_only=True)
    ws = wb['Detailed Listing of Outstanding']
    rows = []
    for row in ws.iter_rows(min_row=22, values_only=False):
        cells = [c.value for c in row]
        vals = cells[2:25] if len(cells) > 24 else cells[2:] + [None] * (23 - max(0, len(cells) - 2))
        if not vals[0]: continue
        rows.append(tuple(str(v).strip() if v is not None else '' for v in vals))
    wb.close()
    with connection.cursor() as cur:
        cur.execute('DELETE FROM wip_accrual')
        execute_values(cur,
            'INSERT INTO wip_accrual (type,branch,dept,charge_code,job,local_ref,wip,accrual,net_total,added,age,debtor_creditor,stat,job_branch,controlling_agent,controlling_customer,management_group,exp_group,orig,eta,etd,posted_by,posted_by_fullname) VALUES %s',
            rows)
    print(f'  -> inserted {len(rows)} wip_accrual rows')

# ========== CREDITOR ==========
cr_matches = [m for m in all_msgs
              if 'creditor transaction report - moc' in m.get('subject', '').lower()
              and m.get('hasAttachments')]
per_group = defaultdict(list)
for m in cr_matches:
    mm = re.search(r'MOC\s+(\w+)\s*$', m.get('subject', '').strip())
    if mm:
        per_group[mm.group(1).upper()].append(m)
print(f'CREDITOR: {len(cr_matches)} total across {len(per_group)} groups')
grand = 0
with connection.cursor() as cur:
    cur.execute('TRUNCATE creditor_transactions')
for group, gmsgs in per_group.items():
    gmsgs.sort(key=lambda m: m['receivedDateTime'], reverse=True)
    try:
        fname, raw_bytes = get_xlsx(gmsgs[0]['id'])
        ws = None; wsx = None
        try:
            wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
            ws = wb.active
            max_row, max_col = ws.max_row, ws.max_column
        except Exception:
            wbx = xlrd.open_workbook(file_contents=raw_bytes)
            wsx = wbx.sheet_by_index(0)
            max_row, max_col = wsx.nrows, wsx.ncols

        def cell(r, c):
            if ws is not None:
                return ws.cell(row=r, column=c).value
            if r - 1 < wsx.nrows and c - 1 < wsx.ncols:
                v = wsx.cell_value(r - 1, c - 1)
                return v if v != '' else None
            return None

        periods = []
        for ci in range(5, max_col + 1):
            v = cell(8, ci)
            if v is not None:
                vs = str(v).strip().split('.')[0]
                if re.match(r'^\d{6}$', vs):
                    periods.append((ci, vs))
        if not periods:
            print(f'  {group}: no periods')
            continue
        cur_branch = ''; rs = []
        for ri in range(9, max_row + 1):
            cb = cell(ri, 2)
            if cb is None: continue
            cb = str(cb).strip()
            if cb.startswith('Organization Branch:'):
                mm = re.search(r'\(([^)]*)\)', cb)
                cur_branch = mm.group(1) if mm else ''
                continue
            if 'Total' in cb or 'Grand' in cb or cb == '':
                continue
            cn = str(cell(ri, 3) or '').strip()
            if not cn: continue
            for ci, period in periods:
                val = cell(ri, ci)
                if val is not None and str(val).strip() != '':
                    try:
                        v = Decimal(str(val).replace(',', ''))
                    except Exception:
                        continue
                    rs.append((cb, cn, period, v, group, cur_branch))
        if rs:
            with connection.cursor() as cur:
                execute_values(cur,
                    'INSERT INTO creditor_transactions (creditor, creditor_name, period, value, creditor_group, branch) VALUES %s',
                    rs)
            print(f'  {group}: {len(rs)} rows')
            grand += len(rs)
    except Exception as e:
        print(f'  {group}: ERR {str(e)[:120]}')
print(f'CREDITOR TOTAL: {grand} rows across {len(per_group)} groups')
