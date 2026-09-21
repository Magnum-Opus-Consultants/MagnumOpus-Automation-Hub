"""Set USEUContact.last_touch to the touchpoint number based on Excel sent dates.

The `last_touch` field holds a TP number (1-10). The Excel (Book1 4.xlsx) only
records 'TP1 Sent On', so the rule is:
    - If Excel's TP1 Sent On has a value -> last_touch = "1"
    - Otherwise -> leave alone (blank stays blank)

Never overwrites an existing non-empty last_touch. Never deletes.

Usage (on server, from automations/):
    ../venv/bin/python3 set_last_touch_from_excel.py          # dry-run (default)
    ../venv/bin/python3 set_last_touch_from_excel.py --apply  # write
"""
import argparse
import os
import sys
from collections import Counter

import django
import openpyxl

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'automations.settings')
django.setup()

from dashboard.models import USEUContact  # noqa: E402


MATCH_FIELDS = [
    ('Org. Name', 'org_name'),
    ('Contact Name', 'contact_name'),
    ('Phone', 'phone'),
    ('Email', 'email'),
]
TP1_SENT_COL = 'TP1 Sent On'


def norm(v):
    return '' if v is None else str(v).strip().lower()


def read_excel(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    rows = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        it = ws.iter_rows(values_only=True)
        try:
            header = next(it)
        except StopIteration:
            continue
        header = [str(h).strip() if h is not None else '' for h in header]
        idx = {h: i for i, h in enumerate(header)}
        required = [c for c, _ in MATCH_FIELDS] + [TP1_SENT_COL]
        missing = [c for c in required if c not in idx]
        if missing:
            print(f'[skip sheet "{sheet_name}"] missing columns: {missing}')
            continue
        for row in it:
            if row is None:
                continue
            rec = {col: (row[idx[col]] if idx[col] < len(row) else None) for col in required}
            rows.append(rec)
    return rows


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--excel', default='Book1 4.xlsx')
    parser.add_argument('--apply', action='store_true')
    args = parser.parse_args()

    if not os.path.exists(args.excel):
        print(f'ERROR: file not found: {args.excel}')
        sys.exit(1)

    print(f'Reading {args.excel} ...')
    rows = read_excel(args.excel)
    print(f'Excel rows: {len(rows)}')

    print('Indexing DB ...')
    db_index = {}
    total_db = 0
    for c in USEUContact.objects.all().iterator():
        total_db += 1
        key = (norm(c.org_name), norm(c.contact_name), norm(c.phone), norm(c.email))
        db_index.setdefault(key, []).append(c)
    print(f'DB rows: {total_db}')

    stats = Counter()
    to_save = {}  # id -> contact
    for rec in rows:
        tp1_sent = rec.get(TP1_SENT_COL)
        if tp1_sent is None or str(tp1_sent).strip() == '':
            stats['excel_no_tp1_sent'] += 1
            continue
        stats['excel_has_tp1_sent'] += 1
        key = tuple(norm(rec[col]) for col, _ in MATCH_FIELDS)
        matches = db_index.get(key, [])
        if not matches:
            stats['no_db_match'] += 1
            continue
        for contact in matches:
            current = (contact.last_touch or '').strip()
            if current != '':
                stats['db_already_set'] += 1
                continue
            stats['would_set_to_1'] += 1
            to_save[contact.id] = contact

    print('\n===== REPORT =====')
    for k, v in stats.items():
        print(f'  {k}: {v}')
    print(f'  unique DB rows to update: {len(to_save)}')

    if not args.apply:
        print('\nDRY-RUN. Re-run with --apply to write.')
        return

    print('\nApplying ...')
    n = 0
    for c in to_save.values():
        c.last_touch = '1'
        c.save(update_fields=['last_touch'])
        n += 1
    print(f'Done. Updated {n} rows.')


if __name__ == '__main__':
    main()
