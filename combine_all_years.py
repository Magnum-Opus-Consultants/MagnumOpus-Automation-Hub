from openpyxl import load_workbook
from datetime import datetime
import os
import psycopg2
from psycopg2.extras import execute_values
import sys

sys.stdout.reconfigure(line_buffering=True)

base_folder = "/home/ethan/Desktop/Gork Test/Turn over GABE tuesday update"

DB_HOST = "167.88.43.168"
DB_PORT = "5432"
DB_NAME = "turnover_data"
DB_USER = "powerbi"
DB_PASSWORD = "your_secure_password"

years = ['2020', '2021', '2022', '2023', '2024', '2025']

def get_branch(filename):
    name = filename.upper().replace('.XLSX', '')
    if 'ALL OVER' in name:
        return None
    return name

def parse_month_code(code):
    try:
        code_str = str(int(code))
        year = int(code_str[:4])
        month = int(code_str[4:])
        if 1 <= month <= 12:
            return datetime(year, month, 1)
    except:
        pass
    return None

def process_wide_format(sheet, branch):
    data = []
    month_columns = {}
    for row in sheet.iter_rows(min_row=1, values_only=True):
        if row and len(row) > 6 and row[3] == 'Debtor':
            for col_idx in range(6, len(row)):
                if row[col_idx]:
                    date_val = parse_month_code(row[col_idx])
                    if date_val:
                        month_columns[col_idx] = date_val
            break

    if not month_columns:
        return data

    for row in sheet.iter_rows(min_row=1, values_only=True):
        if not row or len(row) < 6:
            continue
        debtor_code = row[3]
        debtor_name = row[4]
        if not debtor_code or debtor_code == 'Debtor':
            continue
        if not isinstance(debtor_code, str):
            continue
        for col_idx, date_val in month_columns.items():
            if col_idx < len(row) and row[col_idx] is not None:
                try:
                    value = float(row[col_idx])
                    if value != 0:
                        data.append((
                            debtor_code,
                            debtor_name,
                            value,
                            date_val.strftime('%Y-%m-%d'),
                            date_val.strftime('%Y-%m-%d'),
                            branch,
                            date_val.strftime('%Y-%m-%d')
                        ))
                except (ValueError, TypeError):
                    continue
    return data

print("=" * 50)
print("STARTING DATA PROCESSING")
print("=" * 50)

all_data = []

for year in years:
    year_folder = os.path.join(base_folder, year)
    if not os.path.exists(year_folder):
        print(f"[ERROR] Folder not found: {year_folder}")
        continue

    files = [f for f in os.listdir(year_folder) if f.lower().endswith('.xlsx')]
    files = [f for f in files if 'ALL OVER' not in f.upper()]

    print(f"\n[YEAR {year}] Processing {len(files)} files...")

    for filename in files:
        branch = get_branch(filename)
        if branch is None:
            continue
        filepath = os.path.join(year_folder, filename)
        try:
            wb = load_workbook(filepath, data_only=True)
            sheet = wb.active
            data = process_wide_format(sheet, branch)
            all_data.extend(data)
            print(f"  -> {filename}: {len(data)} rows")
            wb.close()
        except Exception as e:
            print(f"  [ERROR] {filename}: {e}")

print(f"\n[TOTAL] {len(all_data)} rows collected")

print(f"\n[POSTGRESQL] Connecting to {DB_HOST}...")
conn = None
try:
    conn = psycopg2.connect(
        host=DB_HOST,
        port=DB_PORT,
        database=DB_NAME,
        user=DB_USER,
        password=DB_PASSWORD,
        connect_timeout=10
    )
    conn.autocommit = False
    cursor = conn.cursor()
    print("[POSTGRESQL] Connected!")

    print("[POSTGRESQL] Dropping old table if exists...")
    cursor.execute("DROP TABLE IF EXISTS turnover_data CASCADE")
    conn.commit()
    print("[POSTGRESQL] Old table dropped!")

    print("[POSTGRESQL] Creating new table...")
    cursor.execute('''
        CREATE TABLE turnover_data (
            id SERIAL PRIMARY KEY,
            debtor VARCHAR(100),
            debtor_name VARCHAR(255),
            value DECIMAL(15,2),
            date DATE,
            date_fixed VARCHAR(20),
            branch VARCHAR(50),
            date_powerbi DATE
        )
    ''')
    conn.commit()
    print("[POSTGRESQL] Table created!")

    # Bulk insert in batches
    batch_size = 1000
    total = len(all_data)
    print(f"[POSTGRESQL] Bulk inserting {total} rows in batches of {batch_size}...")

    for i in range(0, total, batch_size):
        batch = all_data[i:i+batch_size]
        execute_values(cursor, '''
            INSERT INTO turnover_data (debtor, debtor_name, value, date, date_fixed, branch, date_powerbi)
            VALUES %s
        ''', batch)
        conn.commit()
        print(f"  -> Inserted {min(i+batch_size, total)}/{total} rows...")

    print(f"[POSTGRESQL] Done! {total} rows inserted.")

    cursor.close()
    conn.close()

except psycopg2.OperationalError as e:
    print(f"[POSTGRESQL ERROR] Connection failed: {e}")
    if conn:
        conn.close()
except Exception as e:
    print(f"[POSTGRESQL ERROR] {e}")
    if conn:
        conn.rollback()
        conn.close()

print("\n" + "=" * 50)
print("COMPLETE!")
print("=" * 50)
print(f"\nPower BI Connection:")
print(f"  Server: {DB_HOST}")
print(f"  Port: {DB_PORT}")
print(f"  Database: {DB_NAME}")
print(f"  Username: {DB_USER}")
print(f"  Table: turnover_data")
